# -*- coding: utf-8 -*-
"""
1) stock_reports_queries.sql  - 31 stock report queries ki ek file
2) REPORTS_LIST.xlsx          - category-wise consolidated Excel
gen_reports.py ko import karta hai (jo txt files bhi regenerate karta hai - idempotent).
"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from sql_queries import SQL_QUERIES, GENERIC_QUERIES
from gen_reports import (category, SUB_MAP, menu_path, safe_fname,
                         STOCK_TABLES, DUMP_TABLES, names)
from report_layouts import layout_for

# openpyxl check
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ============ 1) stock_reports_queries.sql ============
NICE = {
 "stockregister":"STOCK REGISTER", "stocksumm":"STOCK SUMMARY",
 "stockregstore":"STORE-WISE STOCK REGISTER", "stocksummstore":"STORE-WISE STOCK SUMMARY",
 "stockinhand":"STOCK IN HAND", "stocksummaryp/sbasis":"STOCK SUMMARY (P/S BASIS)",
 "kitchenstkrep":"KITCHEN STOCK REPORT", "kitchenstksumm":"KITCHEN STOCK SUMMARY",
 "excessconsumption":"EXCESS CONSUMPTION", "restissue":"RESTAURANT-WISE ISSUE",
 "storeissreg":"STORE ISSUE REGISTER", "storeissuereport":"STORE ISSUE REPORT",
 "dailystoreissrpt":"DAILY STORE ISSUE", "issuereg":"ISSUE REGISTER",
 "issueregister":"ISSUE REGISTER (DEPARTMENT-WISE)", "abcanalysis":"ABC ANALYSIS",
 "abcanalysissale":"ABC ANALYSIS (SALE BASIS)", "purchbill":"PURCHASE BILL REPORT",
 "purchorder":"PURCHASE ORDER REPORT", "purchasereg":"PURCHASE REGISTER",
 "purchasesumm":"PURCHASE SUMMARY (PARTY-WISE)", "purchaseledger":"PURCHASE LEDGER",
 "cashcreditpurch":"CASH/CREDIT PURCHASE", "indentreg":"INDENT REGISTER",
 "foodcost":"FOOD COST", "fbcoststatement":"F&B COST STATEMENT",
 "productionreport":"PRODUCTION REPORT", "liquorsalerep":"LIQUOR SALE REPORT",
 "costanalysis":"COST ANALYSIS",
}

bar = "=" * 78
with open("stock_reports_queries.sql", "w", encoding="utf-8") as f:
    f.write(bar + "\n")
    f.write(" STOCK / INVENTORY REPORTS - READY-TO-RUN SQL QUERIES (31)\n")
    f.write(" Database : visahl.sql dump (SQL Server / T-SQL)\n")
    f.write(" Tables   : " + DUMP_TABLES + "\n")
    f.write(" NOTE     : Dump me sirf ye 5 tables hain (ItemMast/PartyMast master tables\n")
    f.write("            dump me NAHI hain). Item names PartyCode ke codes me aayenge.\n")
    f.write("            Dates example hain (@F/@T) - apni dates daal do.\n")
    f.write(bar + "\n\n")
    for key, q in SQL_QUERIES.items():
        title = NICE.get(key, key.upper())
        f.write("-- " + "-" * 76 + "\n")
        f.write(f"-- {title}   (report: {key})\n")
        f.write("-- " + "-" * 76 + "\n")
        f.write(q.strip() + "\n\n")
    f.write("\n" + bar + "\n")
    f.write(" APPENDIX: NON-STOCK CATEGORIES KE GENERIC QUERIES\n")
    f.write(bar + "\n\n")
    for cat, q in GENERIC_QUERIES.items():
        f.write(f"-- [{cat}]\n{q.strip()}\n\n")
print("stock_reports_queries.sql written")

# ============ 2) REPORTS_LIST.xlsx ============
inv_reports = [n for n in names if category(n) == "STOCK / INVENTORY"]

hdr_fill = PatternFill("solid", fgColor="1F4E78")
hdr_font = Font(bold=True, color="FFFFFF", size=11)
stock_fill = PatternFill("solid", fgColor="FFF2CC")
wrap = Alignment(wrap_text=True, vertical="top")

def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

wb = Workbook()

# --- Sheet 1: All Reports ---
ws = wb.active; ws.title = "All Reports"
ws.append(["S.No", "Report Name", "Category", "Stock/Inventory?", "Menu Path", "Source Tables", "Expected Output Columns", "File Name"])
for i, n in enumerate(names, 1):
    cat = category(n); is_inv = cat == "STOCK / INVENTORY"
    subs = SUB_MAP.get(n.lower(), [])
    ws.append([i, n, cat, "HAAN" if is_inv else "NAHI",
               menu_path(subs), STOCK_TABLES.get(n.lower(), ("",))[0] if is_inv else "",
               layout_for(n, cat), safe_fname(n)])
    if is_inv:
        for c in ws[ws.max_row]:
            c.fill = stock_fill
style_sheet(ws, [7, 34, 24, 16, 46, 34, 52, 32])

# --- Sheet 2: Stock-Inventory (31) ---
ws2 = wb.create_sheet("Stock-Inventory (31)")
ws2.append(["S.No", "Report Name", "Menu Path", "Source Tables", "Query in SQL File?", "Stored Procedure", "Expected Output Columns"])
from sql_queries import SQL_QUERIES as SQ
from build_procs_and_data import PROC_NAME
for i, n in enumerate(inv_reports, 1):
    subs = SUB_MAP.get(n.lower(), [])
    has_q = "YES" if n.lower() in SQ else "-"
    ws2.append([i, n, menu_path(subs), STOCK_TABLES.get(n.lower(), ("Stock + ItemMast",""))[0], has_q,
                PROC_NAME.get(n.lower(), "-"), layout_for(n, "STOCK / INVENTORY")])
    for c in ws2[ws2.max_row]:
        c.fill = stock_fill
style_sheet(ws2, [7, 34, 46, 40, 20, 30, 52])

# --- Sheet 3+: category-wise ---
cats = {}
for n in names:
    cats.setdefault(category(n), []).append(n)
for cat in sorted(cats):
    safe_cat = cat.replace(" / ", "-").replace("/", "-")[:28]
    wsc = wb.create_sheet(safe_cat)  # Excel sheet-name limit 31 chars
    wsc.append(["S.No", "Report Name", "Stock/Inventory?", "Menu Path", "Expected Output Columns", "File Name"])
    for i, n in enumerate(cats[cat], 1):
        is_inv = cat == "STOCK / INVENTORY"
        subs = SUB_MAP.get(n.lower(), [])
        wsc.append([i, n, "HAAN" if is_inv else "", menu_path(subs), layout_for(n, cat), safe_fname(n)])
        if is_inv:
            for c in wsc[wsc.max_row]:
                c.fill = stock_fill
    style_sheet(wsc, [7, 34, 16, 46, 52, 32])

# --- Sheet: Report Layouts (saare layouts ek jagah) ---
wsr = wb.create_sheet("Report Layouts")
wsr.append(["S.No", "Report Name", "Category", "Expected Output Columns"])
for i, n in enumerate(names, 1):
    cat = category(n)
    wsr.append([i, n, cat, layout_for(n, cat)])
    wsr.cell(row=wsr.max_row, column=4).alignment = wrap
style_sheet(wsr, [7, 34, 24, 90])

wb.save("REPORTS_LIST.xlsx")
print("REPORTS_LIST.xlsx written")
print(f"Sheets: {wb.sheetnames}")
