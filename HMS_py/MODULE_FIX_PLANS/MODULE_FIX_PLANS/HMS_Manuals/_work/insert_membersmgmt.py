# -*- coding: utf-8 -*-
"""Insert missing Card_Refund.png into both MembersMgmt manuals (section 3.3)."""
import os, shutil, zipfile
import docx
from docx.shared import Inches, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = r"C:/Users/PC/Downloads/189/113/MANUAL/HMS_Manuals/MembersMgmt"
IMG = os.path.join(BASE, "screenshots", "02_SmartCard", "04_Card_Refund.png")
DOCS = ["MembersMgmt_Module_Training_Manual.docx",
        "MembersMgmt_Module_Training_Manual_English.docx"]

BAK = os.path.join(BASE, "_backup")
os.makedirs(BAK, exist_ok=True)

def add_caption(d, text):
    cp = d.add_paragraph()
    r = cp.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return cp

for fn in DOCS:
    src = os.path.join(BASE, fn)
    shutil.copy(src, os.path.join(BAK, fn.replace(".docx", "_pre_cardrefund.docx")))
    d = docx.Document(src)
    s = d.sections[0]
    usable = s.page_width - s.left_margin - s.right_margin
    width = Emu(min(int(usable), int(Inches(6.2))))

    paras = list(d.paragraphs)
    heads = [(i, p) for i, p in enumerate(paras)
             if p.style is not None and p.style.name.startswith("Heading") and p.text.strip()]
    idx = next(i for i, p in heads if p.text.strip().startswith("3.3"))
    end = next((i for i, p in heads if i > idx), len(paras))
    seg = paras[idx:end]
    anchor = None
    for p in seg:
        if p.text.strip().startswith("Navigation Path"):
            anchor = p
    if anchor is None:
        anchor = next(p for p in reversed(seg) if p.text.strip())
    print(fn, "-> section:", seg[0].text.strip(), "| anchor:", anchor.text.strip()[:60])

    tmp = d.add_paragraph()
    tmp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tmp.add_run().add_picture(IMG, width=width)
    anchor._p.addnext(tmp._p)
    cap = add_caption(d, "Screenshot: Card Refund")
    tmp._p.addnext(cap._p)

    d.save(src)
    z = zipfile.ZipFile(src)
    assert z.testzip() is None
    z.close()
    d2 = docx.Document(src)
    n = sum(1 for p in d2.paragraphs if p._p.find(".//" + qn("a:blip")) is not None)
    print("   saved, blips now:", n)

print("DONE")
