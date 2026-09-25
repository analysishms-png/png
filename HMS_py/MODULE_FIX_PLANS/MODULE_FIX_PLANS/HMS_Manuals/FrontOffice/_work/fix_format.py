# -*- coding: utf-8 -*-
"""Fix formatting: center images, add missing captions, scale overflow images, drop empty headings."""
import re, shutil
import docx
from docx.shared import Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = r"C:/Users/PC/Downloads/189/113/MANUAL/HMS_Manuals/FrontOffice"
SRC = BASE + "/FrontOfficeModuleManual.docx"
BAK = BASE + "/_backup/FrontOfficeModuleManual_pre_format_fix.docx"
shutil.copy(SRC, BAK)
print("backup ->", BAK)

d = docx.Document(SRC)
sec = d.sections[0]
usable = sec.page_width - sec.left_margin - sec.right_margin
print("page usable width:", round(usable / 914400, 2), "in")

def clean(s):
    s = re.sub(r"^SECTION\s+\d+\s*\|\s*", "", s.strip())
    s = re.sub(r"^\d+\.\d+\s+", "", s)
    return s.strip()

def has_img(p):
    return p._p.find(".//" + qn("a:blip")) is not None

def is_centered(p):
    jc = p._p.find(qn("w:pPr") + "/" + qn("w:jc"))
    return jc is not None and jc.get(qn("w:val")) == "center"

def next_p_elem(el):
    n = el.getnext()
    return n if n is not None and n.tag == qn("w:p") else None

def text_of(el):
    return "".join(el.itertext()).strip()

def add_caption(text):
    cp = d.add_paragraph()
    r = cp.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return cp

# ---- 1. remove empty heading paragraphs ----
removed = 0
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and not p.text.strip() and not has_img(p):
        p._p.getparent().remove(p._p)
        removed += 1
print("empty headings removed:", removed)

# ---- 2. scale overflowing images ----
scaled = 0
for shp in d.inline_shapes:
    if shp.width > usable:
        ratio = usable / shp.width
        shp.height = Emu(int(shp.height * ratio))
        shp.width = Emu(int(usable))
        scaled += 1
# anchored (floating) images: scale their extent + spPr ext manually
for anc in d.element.body.findall(".//" + qn("wp:anchor")):
    ext = anc.find(qn("wp:extent"))
    aext = anc.find(".//" + qn("a:xfrm") + "/" + qn("a:ext"))
    if ext is not None and int(ext.get("cx")) > usable:
        ratio = usable / int(ext.get("cx"))
        ext.set("cx", str(int(usable)))
        ext.set("cy", str(int(int(ext.get("cy")) * ratio)))
        if aext is not None:
            aext.set("cx", str(int(usable)))
            aext.set("cy", str(int(int(aext.get("cy")) * ratio)))
        scaled += 1
print("images scaled to fit:", scaled)

# ---- 3. center pure-image paragraphs + add missing captions ----
paras = d.paragraphs
first_heading = False
cur = "(cover)"
caps = 0
centered_n = 0
per_sec = {}

for p in paras:
    if p.style.name.startswith("Heading") and p.text.strip():
        cur = clean(p.text)
        first_heading = True
        continue
    if has_img(p):
        if not p.text.strip() and not is_centered(p):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            centered_n += 1
        if not first_heading:
            continue  # cover banner - no caption
        nx = next_p_elem(p._p)
        if nx is not None and text_of(nx).startswith("Screenshot:"):
            continue  # already captioned
        n = per_sec.get(cur, 0) + 1
        per_sec[cur] = n
        label = cur + (" (contd.)" if n > 1 else "")
        cp = add_caption("Screenshot: " + label)
        p._p.addnext(cp._p)
        caps += 1

print("centered:", centered_n)
print("captions added:", caps)

d.save(SRC)
print("saved.")

# ---- 4. verification ----
import zipfile
z = zipfile.ZipFile(SRC)
assert z.testzip() is None
z.close()

d2 = docx.Document(SRC)
nocap = 0
uncentered = 0
wide = 0
first_heading = False
for p in d2.paragraphs:
    if p.style.name.startswith("Heading") and p.text.strip():
        first_heading = True
    if has_img(p):
        nx = p._p.getnext()
        if first_heading and not (nx is not None and text_of(nx).startswith("Screenshot:")):
            nocap += 1
        if not p.text.strip() and not is_centered(p):
            uncentered += 1
for shp in d2.inline_shapes:
    if shp.width > usable:
        wide += 1
print("VERIFY -> body images without caption:", nocap, "| uncentered:", uncentered, "| overflow:", wide)
print("total captions:", sum(1 for p in d2.paragraphs if p.text.strip().startswith("Screenshot:")))
print("DONE")
