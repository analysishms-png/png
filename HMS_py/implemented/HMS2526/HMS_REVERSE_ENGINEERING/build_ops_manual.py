"""HMS_Operations_Manual.docx — Frontend/Backend/DB/SQL workflows + screenshots."""
import os, glob
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
SS = os.path.join(HERE, "screenshots")

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)

FIG = [0]
def FIGURE(path, caption):
    if not path or not os.path.exists(path): return
    FIG[0] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(5.9))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run("Figure %d — %s" % (FIG[0], caption)); r.bold = True; r.font.size = Pt(9)

def H(t, lvl=1): doc.add_heading(t, level=lvl)
def P(t, bold=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold; return p
def B(t): doc.add_paragraph(t, style="List Bullet")
def CODE(t):
    p = doc.add_paragraph(); r = p.add_run(t)
    r.font.name = "Consolas"; r.font.size = Pt(8)

def first(folder, pattern):
    fs = sorted(glob.glob(os.path.join(SS, folder, pattern)))
    return fs[0] if fs else None

H("HMS OPERATIONS MANUAL", 0)
P("Frontend → Backend → Database → SQL Queries → Workflows (Master + Screen Operations)", bold=True)
P("HMS.exe (VB6) + MOONData2627 (SQL Server 2008 R2) | Site 'KK' | 2026-09-28")
P("Har section evidence-backed: [V-VB6] decompiled code, [V-SQL] live trace/catalog, "
  "[V-UI] live GUI capture, [V-CFG] config. Text version: HMS_OPERATIONS_MANUAL.md")

H("1. Frontend Shell", 1)
B("MDI frame ThunderRT6MDIForm; left strip 11 module buttons; top bar per-module menu buttons (TopBarLib).")
B("Child forms 'Module > Group > Screen'; status bar clocks; menu = 524 hardcoded + MenuHelp per-user override.")
FIGURE(first("02_Main_Setup", "014_*.png"), "Main dashboard (MDI) after login")

H("2. Login Workflow", 1)
CODE("userMast COUNT -> USERMAST select -> PASSWD decrypt+UCase compare -> ActiveYN\n"
     "-> AllowDtChng -> User1 JOIN User_Module (AEDP) -> company grid -> comp_code='2'\n"
     "-> Enviro(LogSite_Code='KK') -> MDI")
FIGURE(first("01_Login_Security", "001_*.png"), "Login screen")
FIGURE(first("01_Login_Security", "004_*.png"), "Invalid User validation (live)")
FIGURE(first("01_Login_Security", "010_*.png"), "Company selection grid")

H("3. Master Operations", 1)
H("3.1 Tax Master / Payment Type / Parameters", 2)
B("TaxStru → RevMast.TaxStru → RoomOcc.RoomTaxStru backfill (startup normalize).")
B("Parameters Enviro columns me (221+).")
FIGURE(first("02_Main_Setup", "C3_S1_*Tax_M*"), "Tax Master")
FIGURE(first("02_Main_Setup", "C3_03_*Payme*"), "Payment Type")
FIGURE(first("02_Main_Setup", "C3_04_*Param*"), "Parameter screen")
H("3.2 Users & Rights", 2)
B("userMast + user1.PARAM_STR 'AEDP' + MenuHelp per-user tree (69 users; SA=440 items vs kot=4).")
FIGURE(first("02_Main_Setup", "C3_05_*Print*"), "Printing Parameters")

H("4. Reservation Workflow", 1)
CODE("Insert Into Booking (56 cols) -> BookingMore -> letters (BOOKINGPrint.rpt)\n"
     "Cancel -> BookingCancelDetails (Advance, CancelAmt, CancellationMode)")
FIGURE(first("04_Reservation", "C3_S1_*Reserva*"), "Reservation / Cancellation")
FIGURE(first("04_Reservation", "C3_04_*Reserva*"), "Reservation With History")
FIGURE(first("04_Reservation", "C3_03_*Look_Up*"), "Look Up Rooms")

H("5. Check-In Workflow", 1)
CODE("WalkIn CheckIn chooser -> GuestFolio (40 cols) + RoomOcc (35 cols)\n"
     "+ GuestFolioProfDetail; RoomMast occupied")
FIGURE(first("05_Front_Office", "C4_WalkInCheckIn_Form.png"), "WalkIn CheckIn chooser (6 actions)")
FIGURE(first("05_Front_Office", "C3_S1_*WalkIn*"), "WalkIn CheckIn form")
FIGURE(first("05_Front_Office", "C3_02_*Room_S*"), "Room Status (live occupancy)")

H("6. Posting, Settlement, Check-Out", 1)
CODE("Charge: PayCharge(AmtDr, TaxPer, OnAmt) -> Payment: PayCharge(AmtCr, SettleDate, Bill_No)\n"
     "Bill: BillDet view = SUM(AmtDr-AmtCr) -> BillPrint.rpt\n"
     "CheckOut: Update RoomOcc set CHKOUTDATE/CHKOUTTIME")
FIGURE(first("05_Front_Office", "C3_03_*Expens*"), "Expense Entry (posting)")
FIGURE(first("05_Front_Office", "C3_05_*Bill_R*"), "Bill Reprint")

H("7. Night Audit (documented, NOT executed)", 1)
CODE("1 KOT/POS pending check -> 2 TRANS: ExpSheet/RoomChrgDueAc -> Ledger -> COMMIT\n"
     "3 dirty rooms roomStat='D' -> 4 due-out extend -> 5 enviro.ncur+1 + NightAuditLog\n"
     "6 tokens reset + split staging delete + prefix renumber -> COMMIT")
FIGURE(first("10_Night_Audit", "C3_S1_*GSTR-1*"), "Night Audit module — GST reports (reports only)")
FIGURE(first("10_Night_Audit", "C2_*Reconcilia*"), "GSTR-2A Reconciliation")

H("8. POS Workflows (3 outlets)", 1)
B("Main + High Way point + MOON TEA CAFE (live menus). KOT -> Sale1/2 -> settle -> room-post/print.")
FIGURE(first("08_POS", "C2_*Sales_Re*"), "POS Sales Register")
FIGURE(first("08_POS", "C3_P5C5_*"), "High Way point — Sale Bill Entry")
FIGURE(first("08_POS", "C3_P5C6_*"), "High Way point — Settlement Entry")
FIGURE(first("08_POS", "C3_P5D8_*"), "MOON TEA CAFE — Bill Lookup")

H("9. Banquet", 1)
FIGURE(first("09_Banquet", "C2_*Banquet_Booking*"), "Banquet Booking")
FIGURE(first("09_Banquet", "C2_*Chef_Pre-Costing*"), "Chef Pre-Costing")
FIGURE(first("09_Banquet", "C3_P5D7_*Bill_Look_Up*"), "Bill Look Up")

H("10. Inventory", 1)
FIGURE(first("07_Inventory", "C3_S1_*Indent*"), "Indent")
FIGURE(first("07_Inventory", "C3_02_*Purchase_*"), "Purchase Order")
FIGURE(first("07_Inventory", "C3_04_*Purchase_*"), "Purchase Bill")

H("11. HR / House Keeping / Messaging", 1)
FIGURE(first("11_HR_Payroll", "C3_S1_*Leave*"), "Leave entry")
FIGURE(first("11_HR_Payroll", "C2_*LoanAdvance*"), "Loan/Advance")
FIGURE(first("06_House_Keeping", "C3_01_Complain_Master*"), "Complain Master")
FIGURE(first("06_House_Keeping", "C2_*House_Kee*"), "House Keeping Screen")
FIGURE(first("13_Messaging", "C3_INBOX_Inbox*"), "In Box (in-app messaging)")

H("12. Finance", 1)
FIGURE(first("03_Finance", "C3_S1_*Voucher_En*"), "Voucher Entry")

H("13. GST e-Invoice", 1)
B("Staging eInvoiceBill1/2 -> JSON v1.1 -> CI GSP -> IRN/AckNo/SignedQRCode (185/211 live) -> printed QR.")

H("14. SQL Quick Reference", 1)
CODE("Business date: SELECT NCUR FROM Enviro WHERE LogSite_Code='KK'\n"
     "In-house: GuestFolio JOIN RoomOcc WHERE CHKOUTDATE IS NULL\n"
     "Bill: SELECT * FROM BillDet WHERE LogSite_Code='KK'\n"
     "Last audits: TOP 5 DateChngTo,U_Name FROM NightAuditLog\n"
     "User menu: MenuHelp WHERE UserName=.. AND CompCode='2'")

H("15. Cautions", 1)
B("Night Audit backup ke bina nahi; duplicate-PK bugs (concurrent posting); LocationMast broken;")
B("credentials DB/plain-text me — backups restricted (SRN-001); shared logins avoid (U_Name audit).")

H("16. Evidence Index", 1)
B("CAPTURE_INDEX.md (97 PNGs/71 windows), MENU_INVENTORY (524), MENUHELP analysis,")
B("CHECKIN_FIELD_MAP, NIGHT_AUDIT_FLOW, LIVE_DATABASE_DICTIONARY, SQL_TRACKING_RESULTS,")
B("GST_EINVOICE_FLOW, BUG_REGISTER, SRN-001, MODERNIZATION_BLUEPRINT.")

doc.save(os.path.join(HERE, "22_Final_Manual", "HMS_Operations_Manual.docx"))
print("Ops manual saved with", FIG[0], "figures")
