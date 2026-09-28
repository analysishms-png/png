"""Build HMS_Modernization_Blueprint.docx with embedded charts."""
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
CH = os.path.join(HERE, "charts")

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)

def H(t, lvl=1):
    doc.add_heading(t, level=lvl)

def P(t, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    return p

def B(t):
    doc.add_paragraph(t, style="List Bullet")

def IMG(fname, w=6.3, cap=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(CH, fname), width=Inches(w))
    if cap:
        c = doc.add_paragraph(cap)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].italic = True
        c.runs[0].font.size = Pt(9)

# Cover
H("HMS Modernization Blueprint", 0)
P("VB6 + SQL Server 2008 R2 -> Phased Modern Stack", bold=True)
P("Date: 2026-09-28 | Rule: business rules preserved 1:1 — naya behavior invent nahi kiya gaya.")
P("Evidence base: decompiled HMS.exe (351 objects), live MOONData2627 catalog (277 tables), "
  "runtime error logs, 524 Crystal Reports. Full technical detail: MODERNIZATION_BLUEPRINT.md")

H("1. Measured Baseline (Charts)", 1)
IMG("chart1_top_tables.png", cap="Figure 1 — Top 15 tables by SQL usage in decompiled code [VERIFIED-VB6]")
P("Core revenue chain: Depart(1378), PayCharge(873), GuestFolio(869), GuestProf(840), RevMast(835). "
  "In tables ke column semantics migration me sabse pehle freeze honge.")
IMG("chart2_module_complexity.png", cap="Figure 2 — VB6 forms/modules per module (351 total) [VERIFIED-VB6]")
P("Front Office (55) aur POS (49) sabse bade UI surfaces hain — inhi ka test coverage sabse deep hoga.")
IMG("chart3_bug_severity.png", cap="Figure 3 — Documented bugs by severity [VERIFIED-SQL/VB6]")
P("9 bugs registered; 2 HIGH (PayCharge duplicate-PK posting crash, hardcoded 'India12' backdoor — SRN-001).")
IMG("chart4_db_objects.png", cap="Figure 4 — MOONData2627 live object counts [VERIFIED-SQL]")
P("277 tables lekin 0 triggers / 0 FKs / 0 functions — saara integrity VB6 app me hai. "
  "Isliye DB-only migration kaafi nahi; service layer me rules re-create honge.")
IMG("chart5_effort.png", cap="Figure 5 — Relative effort estimate [INFERRED — code volume + coupling se derived]")

H("2. Target Stack", 1)
tbl = doc.add_table(rows=1, cols=2)
tbl.style = "Light Grid Accent 1"
hdr = tbl.rows[0].cells
hdr[0].text = "Layer"
hdr[1].text = "Target"
rows = [
 ("UI", "Web app (React/Vue ya Blazor) — hotel LAN + optional cloud"),
 ("Business logic", "Service layer (FastAPI/.NET) — same rules, testable"),
 ("Data access", "Repository + server-side transactions (BUG-007 class fix)"),
 ("Database", "SQL Server 2019+ — schema pehle 1:1 compatible"),
 ("Reports", "Modern engine; .ttx field maps reuse"),
 ("Auth", "Central auth — bcrypt/argon2, lockout, claims (SRN-001)"),
 ("Side stores", "Jet .mdb (DMEPABX/Fadata/Temp/SaleTransfer) -> SQL"),
 ("Config", "Analysis.ini -> config service + secrets vault"),
]
for a, b in rows:
    r = tbl.add_row().cells
    r[0].text = a
    r[1].text = b

H("3. Phased Roadmap", 1)
B("Phase A — Foundation (4–6 wks): SQL upgrade, schema snapshot 1:1, auth service, CI/CD + test DB.")
B("Phase B — Core pilot (8–12 wks): FO read-only screens, Reservation/Check-In dual-run, Night Audit service on test DB.")
B("Phase C — Transactions (12–16 wks): Posting/Payment/Check-Out, POS+KOT, Finance vouchers + Tally export.")
B("Phase D — Cutover (8–12 wks): remaining modules, reports replacement, VB6 retire, Jet stores migrate.")
P("Dual-run principle: purana aur naya output side-by-side verify (PayCharge rows, bill totals, "
  "NightAuditLog) — 100% match ke bina cutover nahi.", bold=True)

H("4. Priority Matrix", 1)
B("P1: Front Office, Night Audit, Reservation (revenue core — chart1 coupling)")
B("P2: Finance/GL, POS, Main Setup/Security")
B("P3: Inventory, Banquet, Members/SmartCard, House Keeping")
B("P4: HR & Payroll (SaaS option), EPABX, Messaging, Mall Mgmt")

H("5. Key Risks & Mitigations", 1)
B("Hidden logic in 351 forms -> decomp + live SQL tracing continue; per-module freeze docs.")
B("0 FK/0 trigger -> app-side rules document karke service layer me re-create.")
B("Night Audit regression = revenue corruption -> exact SQL port + test-DB parallel audit + log diff.")
B("Counter desync (BUG-003/004) -> centralized numbering service.")
B("Report parity -> 524 .rpt me se priority wale report-by-report signoff (.ttx maps ready).")

H("6. Timeline & Next Steps", 1)
P("Realistic total: 8–12 months phased (3–5 team members); P1 modules pehle live, VB6 parallel chalta rahega.")
B("Owner sign-off: target stack + Phase A budget")
B("SRN-001 containment (password rotation) — immediate")
B("GUI walkthrough complete -> module manuals final")
B("Front Office field-level spec freeze (screens + SQL + validations)")

doc.save(os.path.join(HERE, "HMS_Modernization_Blueprint.docx"))
print("DOCX saved")
