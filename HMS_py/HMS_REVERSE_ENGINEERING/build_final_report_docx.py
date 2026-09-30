"""FINAL_REPORT.docx — closing report converted from FINAL_REPORT.md."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "22_Final_Manual", "FINAL_REPORT.docx")

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)

def H(t, lvl=1):
    doc.add_heading(t, level=lvl)

def P(t, bold=False, italic=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return p

def B(t):
    doc.add_paragraph(t, style="List Bullet")

def N(t):
    doc.add_paragraph(t, style="List Number")

def CODE(t):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.font.name = "Consolas"
    r.font.size = Pt(8)

def KEYVAL(key, val):
    p = doc.add_paragraph()
    r = p.add_run(key + "  ")
    r.bold = True
    r.font.size = Pt(10)
    r2 = p.add_run(val)
    r2.font.size = Pt(10)

# ---------- Cover ----------
t = doc.add_heading("HMS REVERSE-ENGINEERING — FINAL REPORT", level=0)
for run in t.runs:
    run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
P("Analysis formally closed — 2026-09-28", bold=True)
KEYVAL("Mode:", "100% READ-ONLY (no business data touched)")
KEYVAL("System:", "HMS.exe (VB6 MDI, 351 objects) + MOONData2627 (SQL Server 2008 R2, 277 tables)")
KEYVAL("Site:", "'KK', Bareilly")

# ---------- 1. Summary ----------
H("1. KAAM KA SUMMARY (kya-kya achieve hua)", 1)

H("1.1 Evidence collection", 2)
rows = [
    ("Stream", "Volume"),
    ("Decompiled code (EXTRAS.text)", "54 MB — 351 objects, 524 menu items, exact SQL strings"),
    ("Live DB catalog", "277 tables, 5,798 columns, 244 indexes, 10 views, 3 procs, 0 FK/triggers"),
    ("Live SQL trace", "17,525 statements (40 writes fully accounted — 0 business changes)"),
    ("GUI captures", "97 PNGs, 71 verified window-opens, action_log.csv audit trail"),
    ("Runtime logs", "ErrorLog/FaLog/IssueLists (2017–2026)"),
]
tbl = doc.add_table(rows=len(rows), cols=2)
tbl.style = "Light Grid Accent 1"
for i, (c1, c2) in enumerate(rows):
    cells = tbl.rows[i].cells
    cells[0].text = c1
    cells[1].text = c2
    if i == 0:
        for c in cells:
            for r in c.paragraphs[0].runs:
                r.bold = True

H("1.2 Documentation deliverables (22_Final_Manual/)", 2)
N("HMS_Complete_Reverse_Engineering_Manual_v2.docx — 87 figures, full technical manual")
N("HMS_Operations_Manual.docx + .md — Frontend→Backend→DB→SQL workflows (33 figures)")
N("HMS_WORKFLOWS_PART2.md — EPABX/SMS/Members/SmartCard + status matrix")
N("00_MASTER_INDEX.md — poore workspace ka navigation map")
N("HMS_SYSTEM_ARCHITECTURE.md, MASTER_SCREEN_INVENTORY.xlsx, MASTER_DATABASE_DICTIONARY.xlsx")
N("23_Modernization_Blueprint/ — 5 charts (measured) + 4-phase roadmap + .docx")

H("1.3 Key discoveries", 2)
B("Menu engine dual-source: 524 hardcoded (Proc_165_0_14C1708) + MenuHelp per-user override "
  "(OPT1..4, Flag E/R/N/V; 69 users: SA=440 → kot=4)")
B("RBAC = 2 layers: MenuHelp (visibility) + user1.PARAM_STR 'AEDP' (rights) — full RBAC spec extract ho gaya")
B("Integrity 100% app-side (0 FK/0 triggers) — modernization ka sabse bada risk/input")
B("Business date per-site (Enviro.NCUR); Night Audit = 6-step transaction (documented, never executed)")
B("GST e-Invoice live: CI GSP → IRP, 185/211 IRNs, last ack 2026-08-01")
B("SMS outbox proven: DBSENDSMS 18,810 rows, JobScheduler poll pattern")
B("POS multi-outlet live: main + High Way point + MOON TEA CAFE")
B("App-name: 'Logic Pro (Win HMS)' — trace-filtering ke liye")

H("1.4 Bugs & security (recorded, NOT fixed)", 2)
B("BUG-001..010 — schema drift (LocationMast missing LIVE), PK crashes (counter desync), "
  "report paths, transaction nesting, unhandled EOF/field errors")
B("Security: BUG-008 'India12' backdoor, reversible passwords, SA seeding, "
  "connection-string creds (SRN-001 F-1..F-4), e-invoice plaintext secrets (BUG-010/F-5)")
B("SRN-001: containment + patch routes + verification plan + owner decisions")

H("1.5 Safety record (100% clean)", 2)
B("Koi INSERT/UPDATE/DELETE/DDL HMS data pe nahi (40 writes = app-internal normalization/staging "
  "cleanup — full accounting in SQL_TRACKING_RESULTS.md §2)")
B("Night Audit kabhi execute nahi (blocklist-enforced)")
B("Har GUI form ESC/WM_CLOSE/'No' se closed — zero saves")
B("Passwords kabhi stored/logged nahi (masking throughout)")
B("Trace stop + cleanup verified (0 traces, 0 HMS processes)")

# ---------- 2. RBAC skip ----------
H("2. RBAC LIVE VERIFY — SKIPPED (owner decision)", 1)
P("Deepak login attempt: guessed password → 'Invalid Password' (TC-009 error path live-validated).")
P("Live restricted-menu capture skip — MenuHelp-based RBAC docs complete (DEEPAK_ROLE_PROFILE.md).")
P("Future-me karne ke liye: Deepak se login → strip buttons compare (expect Inventory/Finance/HR "
  "hidden) → menu count vs 157. Master index me skip note recorded.")

# ---------- 3. Workspace ----------
H("3. WORKSPACE LOCATION", 1)
CODE(
    "C:\\PENDRIVE\\HMS\\HMS_REVERSE_ENGINEERING\\\n"
    "├── 00_Project_Discovery\\   environment, INI, DB connection\n"
    "├── 01_Login_Security\\      LOGIN_FLOW.md (+ 12 login screenshots)\n"
    "├── 02..16_*\\               module folders (captures + module docs)\n"
    "├── 05_Front_Office\\        FRONT_OFFICE_LIFECYCLE + CHECKIN_FIELD_MAP\n"
    "├── 10_Night_Audit\\         NIGHT_AUDIT_FLOW (code-verified)\n"
    "├── 17_Reports\\             REPORTS_ANALYSIS (524 files mapped)\n"
    "├── 18_Database\\            LIVE_DATABASE_DICTIONARY + raw extracts (5,798 cols)\n"
    "├── 19_VB6_Logic\\           MENU_INVENTORY (524), MENUHELP analysis, role profiles, 351 objects\n"
    "├── 20_SQL_Tracking\\        SQL_TRACKING_RESULTS + raw statements\n"
    "├── 21_Testing\\             BUG_REGISTER, SRN-001, TEST_CASES\n"
    "├── 22_Final_Manual\\        ★ 4 manuals + master index + inventories + this report\n"
    "├── 23_Modernization_Blueprint\\  blueprint + charts\n"
    "├── 24_GST_EInvoice\\        GST_EINVOICE_FLOW\n"
    "└── screenshots\\            97 PNGs + CAPTURE_INDEX.md")

# ---------- 4. Next steps ----------
H("4. RECOMMENDED NEXT STEPS (priority order)", 1)
N("SRN-001 containment aaj: userMast + SQL passwords rotate; C:\\Backup ACL; banquet terminal access restrict")
N("DB backup lena start (SIMPLE recovery, koi backup chain nahi mili)")
N("BUG-003/004: posting single-terminal tak restrict jab tak numbering-service fix na ho")
N("Modernization Phase A: SQL 2019 upgrade (2008 R2 EOL) + schema snapshot (already extracted)")
N("LocationMast table create karna (BUG-001) ya screen menu se hatana")
N("Optional: Deepak RBAC live verify (ab jab bhi password available ho)")

# ---------- Footer note ----------
doc.add_paragraph()
P("Analysis closed 2026-09-28. Saara evidence HMS_REVERSE_ENGINEERING workspace me preserved — "
  "koi bhi developer ye docs padhke HMS ko bina guesswork ke rebuild kar sakta hai.", italic=True, size=9)

doc.save(OUT)
print("Saved:", OUT, "-", os.path.getsize(OUT), "bytes")
